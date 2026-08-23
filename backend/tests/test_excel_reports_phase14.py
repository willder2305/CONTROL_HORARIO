from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import create_app, db
from app.models import Administrador, HorarioTrabajador, Marcacion, Trabajador


class TestConfig:
    SECRET_KEY = "test-secret"
    SQLALCHEMY_DATABASE_URI = "sqlite:///:memory:"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    CORS_ORIGINS = ["http://localhost:5173"]
    SESSION_COOKIE_HTTPONLY = True
    SESSION_COOKIE_SAMESITE = "Lax"
    SESSION_COOKIE_SECURE = False
    PERMANENT_SESSION_LIFETIME = 3600
    FINGERPRINT_PROVIDER = "mock"
    TESTING = True


def build_client_with_excel_report_data():
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        db.session.add(
            Administrador(
                usuario="admin",
                password_hash=generate_password_hash("Correcta123"),
                nombre="Admin",
                apellido="Principal",
                activo=True,
            )
        )
        worker_a = Trabajador(codigo="EMP-0001", nombre="Ana", apellido="Lopez")
        worker_b = Trabajador(codigo="EMP-0002", nombre="Luis", apellido="Mendez")
        db.session.add_all([worker_a, worker_b])
        db.session.flush()
        schedule_a = HorarioTrabajador(
            trabajador_id=worker_a.id,
            hora_entrada=time(8, 0),
            hora_salida_almuerzo=time(12, 45),
            hora_regreso_almuerzo=time(13, 45),
            hora_salida=time(18, 0),
            tolerancia_entrada=0,
            tolerancia_regreso_almuerzo=0,
            fecha_inicio=date(2026, 8, 1),
            activo=True,
        )
        schedule_b = HorarioTrabajador(
            trabajador_id=worker_b.id,
            hora_entrada=time(9, 0),
            hora_salida_almuerzo=time(13, 45),
            hora_regreso_almuerzo=time(14, 45),
            hora_salida=time(18, 0),
            tolerancia_entrada=0,
            tolerancia_regreso_almuerzo=0,
            fecha_inicio=date(2026, 8, 1),
            activo=True,
        )
        db.session.add_all([schedule_a, schedule_b])
        db.session.flush()
        db.session.add_all(
            [
                Marcacion(
                    trabajador_id=worker_a.id,
                    fecha=date(2026, 8, 20),
                    fecha_hora=datetime(2026, 8, 20, 8, 10),
                    tipo_marcacion="ENTRADA",
                    hora_programada=time(8, 0),
                    hora_real=time(8, 10),
                    estado="TARDANZA",
                    minutos_diferencia=10,
                    horario_trabajador_id=schedule_a.id,
                    origen="TEST",
                ),
                Marcacion(
                    trabajador_id=worker_a.id,
                    fecha=date(2026, 8, 20),
                    fecha_hora=datetime(2026, 8, 20, 17, 50),
                    tipo_marcacion="SALIDA",
                    hora_programada=time(18, 0),
                    hora_real=time(17, 50),
                    estado="SALIDA_ANTICIPADA",
                    minutos_diferencia=10,
                    horario_trabajador_id=schedule_a.id,
                    origen="TEST",
                ),
                Marcacion(
                    trabajador_id=worker_b.id,
                    fecha=date(2026, 8, 21),
                    fecha_hora=datetime(2026, 8, 21, 9, 0),
                    tipo_marcacion="ENTRADA",
                    hora_programada=time(9, 0),
                    hora_real=time(9, 0),
                    estado="A_TIEMPO",
                    minutos_diferencia=0,
                    horario_trabajador_id=schedule_b.id,
                    origen="TEST",
                ),
            ]
        )
        db.session.commit()

    client = app.test_client()
    client.post("/api/auth/login", json={"usuario": "admin", "password": "Correcta123"})
    return app, client


def read_workbook(response):
    return load_workbook(BytesIO(response.data))


def test_excel_report_requires_session():
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()

    response = app.test_client().get("/api/admin/reportes/excel")

    assert response.status_code == 401


def test_excel_report_downloads_expected_columns_and_rows():
    app, client = build_client_with_excel_report_data()

    response = client.get("/api/admin/reportes/excel")
    workbook = read_workbook(response)
    worksheet = workbook.active

    assert response.status_code == 200
    assert response.content_type.startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert [cell.value for cell in worksheet[1]] == [
        "Fecha",
        "Código",
        "Trabajador",
        "Entrada/Tipo",
        "Hora programada",
        "Hora registrada",
        "Estado",
        "Diferencia",
    ]
    assert worksheet.max_row == 4
    assert worksheet["B2"].value == "EMP-0001"
    assert worksheet["C2"].value == "Ana Lopez"
    assert worksheet["G2"].value == "TARDANZA"


def test_excel_report_uses_selected_filters():
    app, client = build_client_with_excel_report_data()

    response = client.get("/api/admin/reportes/excel?estado=TARDANZA")
    worksheet = read_workbook(response).active

    assert response.status_code == 200
    assert worksheet.max_row == 2
    assert worksheet["G2"].value == "TARDANZA"
