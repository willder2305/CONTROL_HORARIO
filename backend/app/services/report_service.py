from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.utils.datetime_utils import obtener_hora_actual


REPORT_HEADERS = [
    "Fecha",
    "Código",
    "Trabajador",
    "Entrada/Tipo",
    "Hora programada",
    "Hora registrada",
    "Estado",
    "Diferencia",
]


def build_excel_report(marks):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Reporte"
    worksheet.append(REPORT_HEADERS)

    for mark in marks:
        worker = mark.trabajador
        worksheet.append(
            [
                mark.fecha,
                worker.codigo if worker else "",
                f"{worker.nombre} {worker.apellido}" if worker else "",
                mark.tipo_marcacion,
                mark.hora_programada,
                mark.hora_real,
                mark.estado,
                mark.minutos_diferencia,
            ]
        )

    apply_simple_report_format(worksheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_pdf_report(marks, filters, summary):
    """Construye el PDF del reporte administrativo con los filtros aplicados."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.35 * inch,
        bottomMargin=0.35 * inch,
        title="Reporte de control de horarios",
        pageCompression=0,
    )

    styles = build_pdf_styles()
    story = [
        Paragraph("FARMACIA", styles["center_title"]),
        Paragraph("REPORTE DE CONTROL DE HORARIOS", styles["center_subtitle"]),
        Spacer(1, 10),
    ]
    story.extend(build_pdf_filter_table(filters, styles))
    story.append(Spacer(1, 10))
    story.extend(build_pdf_summary_table(summary, styles))
    story.append(Spacer(1, 10))
    story.append(build_pdf_marks_table(marks, styles))

    document.build(story, onFirstPage=draw_page_footer, onLaterPages=draw_page_footer)
    output.seek(0)
    return output


def build_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="center_title",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#1F4E78"),
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="center_subtitle",
            parent=styles["Heading2"],
            alignment=TA_CENTER,
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#263238"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="table_cell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=8.5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="table_header",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.4,
            leading=8.8,
            textColor=colors.white,
        )
    )
    return styles


def build_pdf_filter_table(filters, styles):
    generated_at = obtener_hora_actual().strftime("%Y-%m-%d %H:%M")
    data = [
        ["Fecha de generacion", generated_at, "Rango", filters.get("rango", "Todos")],
        ["Trabajador", filters.get("trabajador", "Todos"), "Tipo", filters.get("tipo", "Todos")],
        ["Estado", filters.get("estado", "Todos"), "", ""],
    ]
    table = Table(data, colWidths=[1.35 * inch, 3.0 * inch, 0.9 * inch, 2.2 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F7FA")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#263238")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2F3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return [table]


def build_pdf_summary_table(summary, styles):
    data = [
        [
            "Total marcaciones",
            "Total tardanzas",
            "Total a tiempo",
            "Salidas anticipadas",
        ],
        [
            summary["total_marcaciones"],
            summary["total_tardanzas"],
            summary["total_a_tiempo"],
            summary["total_salidas_anticipadas"],
        ],
    ]
    table = Table(data, colWidths=[1.8 * inch, 1.8 * inch, 1.8 * inch, 1.8 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2F3")),
                ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#FFFFFF")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return [table]


def build_pdf_marks_table(marks, styles):
    headers = [
        "Fecha",
        "Codigo",
        "Trabajador",
        "Entrada/Tipo",
        "Programada",
        "Registrada",
        "Estado",
        "Diferencia",
    ]
    data = [[Paragraph(header, styles["table_header"]) for header in headers]]

    for mark in marks:
        worker = mark.trabajador
        data.append(
            [
                format_date(mark.fecha),
                worker.codigo if worker else "",
                Paragraph(f"{worker.nombre} {worker.apellido}" if worker else "", styles["table_cell"]),
                mark.tipo_marcacion,
                format_time(mark.hora_programada),
                format_time(mark.hora_real),
                mark.estado,
                f"{mark.minutos_diferencia} min",
            ]
        )

    if len(data) == 1:
        data.append(["Sin resultados", "", "", "", "", "", "", ""])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            0.78 * inch,
            0.88 * inch,
            1.55 * inch,
            1.25 * inch,
            0.86 * inch,
            0.86 * inch,
            1.18 * inch,
            0.82 * inch,
        ],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.2),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def draw_page_footer(canvas, document):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#607D8B"))
    canvas.drawRightString(
        document.pagesize[0] - document.rightMargin,
        0.2 * inch,
        f"Pagina {document.page}",
    )
    canvas.restoreState()


def format_date(value):
    return value.strftime("%Y-%m-%d") if value else ""


def format_time(value):
    return value.strftime("%H:%M") if value else ""


def apply_simple_report_format(worksheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"

    for column in ("E", "F"):
        for cell in worksheet[column][1:]:
            cell.number_format = "hh:mm"

    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 16
    worksheet.column_dimensions["C"].width = 28
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18
    worksheet.column_dimensions["G"].width = 22
    worksheet.column_dimensions["H"].width = 14
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:H{max(worksheet.max_row, 1)}"
